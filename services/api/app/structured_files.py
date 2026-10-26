import csv
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path
import zipfile

from openpyxl import load_workbook


MAX_ROWS = 10_000
MAX_COLUMNS = 200
MAX_SHEETS = 20


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def _headers(row: list[str]) -> list[str]:
    result=[]
    used=set()
    for index, value in enumerate(row, start=1):
        name=value.strip() or f"column_{index}"
        if name in used:
            name=f"{name}_{index}"
        used.add(name);result.append(name)
    return result


def _package_result(sheets: list[tuple[str, list[list[str]]]]) -> dict:
    lines=[]; all_text=[]; preview=[]; row_total=0
    for sheet_index, (sheet_name, rows) in enumerate(sheets, start=1):
        if not rows: continue
        headers=_headers(rows[0][:MAX_COLUMNS])
        for row_number, row in enumerate(rows[1:MAX_ROWS+1], start=2):
            values=[_cell_text(v) for v in row[:MAX_COLUMNS]]
            record={headers[i]:values[i] if i < len(values) else "" for i in range(len(headers))}
            if not any(record.values()): continue
            row_total+=1
            rendered=" | ".join(f"{k}: {v}" for k,v in record.items() if v)
            all_text.append(rendered)
            lines.append({"text":rendered,"confidence":1.0,"page":sheet_index,
                          "coordinate_space":"sheet_cells","source":"STRUCTURED_FILE",
                          "sheet":sheet_name,"row":row_number,"cells":record})
            if len(preview)<5:preview.append(record)
    return {"full_text":"\n".join(all_text),"lines":lines,
            "fields":{"row_count":{"raw_value":row_total,"normalized_value":row_total,"confidence":1.0,"status":"PARSED"},
                      "columns":{"raw_value":list(preview[0]) if preview else [],"normalized_value":list(preview[0]) if preview else [],"confidence":1.0,"status":"PARSED"},
                      "preview_rows":{"raw_value":preview,"normalized_value":preview,"confidence":1.0,"status":"PARSED"}},
            "document_type":"STRUCTURED_DATA","provider":"CSV/XLSX parser","model_version":"structured-rows-v1"}


def extract_csv(data: bytes) -> dict:
    try:
        content=data.decode("utf-8-sig")
    except UnicodeDecodeError:
        content=data.decode("gb18030")
    if "\x00" in content:
        raise ValueError("CSV 包含二进制空字符")
    try:
        dialect=csv.Sniffer().sniff(content[:8192], delimiters=",\t;")
    except csv.Error:
        dialect=csv.excel
    rows=[]
    for index,row in enumerate(csv.reader(StringIO(content),dialect),start=1):
        if len(row)>MAX_COLUMNS:raise ValueError(f"CSV 列数不能超过 {MAX_COLUMNS}")
        rows.append(row)
        if index>MAX_ROWS+1:raise ValueError(f"CSV 行数不能超过 {MAX_ROWS}")
    if not rows:raise ValueError("CSV 没有数据")
    return _package_result([(Path("upload.csv").stem,rows)])


def extract_xlsx(data: bytes) -> dict:
    with zipfile.ZipFile(BytesIO(data)) as archive:
        infos=archive.infolist()
        if len(infos)>2000 or sum(x.file_size for x in infos)>64*1024*1024:
            raise ValueError("XLSX 解压后内容超过安全限制")
        if any(x.file_size/max(1,x.compress_size)>100 for x in infos):
            raise ValueError("XLSX 压缩率超过安全限制")
        names=set(archive.namelist())
        if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
            raise ValueError("不是有效的 XLSX 工作簿")
    workbook=load_workbook(BytesIO(data),read_only=True,data_only=True)
    try:
        if len(workbook.worksheets)>MAX_SHEETS:raise ValueError(f"工作表数量不能超过 {MAX_SHEETS}")
        sheets=[]
        for sheet in workbook.worksheets:
            rows=[]
            for index,row in enumerate(sheet.iter_rows(values_only=True),start=1):
                values=[_cell_text(v) for v in row]
                while values and not values[-1]:values.pop()
                if len(values)>MAX_COLUMNS:raise ValueError(f"工作表 {sheet.title} 列数超过 {MAX_COLUMNS}")
                rows.append(values)
                if index>MAX_ROWS+1:raise ValueError(f"工作表 {sheet.title} 行数超过 {MAX_ROWS}")
            sheets.append((sheet.title,rows))
        return _package_result(sheets)
    finally:
        workbook.close()
